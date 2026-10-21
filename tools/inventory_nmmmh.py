#!/usr/bin/env python3
"""Inventory available original documents, without pretending they are audited."""
import argparse
import hashlib
import json
import re
from pathlib import Path

import pymupdf
from docx import Document

WEB = Path(__file__).resolve().parents[1]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', type=Path, required=True)
    args = parser.parse_args()
    root = args.source.resolve()
    manifest = json.loads((WEB/'subjects/nmmmh/manifest.json').read_text())
    old = json.loads((root/'DAPAN_KIEMCHUNG/exams.json').read_text())
    published = {e['id'] for e in manifest['exams']}
    documents = []
    for file in sorted(root.rglob('*')):
        if not file.is_file() or file.suffix.lower() not in ('.pdf', '.docx'):
            continue
        relative = str(file.relative_to(root))
        record = {'path': relative, 'sha256': hashlib.sha256(file.read_bytes()).hexdigest(), 'bytes': file.stat().st_size}
        if file.suffix.lower() == '.pdf':
            with pymupdf.open(file) as doc:
                texts = [p.get_text() for p in doc]
                record.update(pages=len(doc), native_text_chars=sum(len(x.strip()) for x in texts), question_header_occurrences=len(re.findall(r'Câu\s*\d+', '\n'.join(texts), re.I)))
                record['strategy'] = 'native_text_first' if record['native_text_chars'] > len(doc)*100 else 'scan_requires_ocr_and_visual_check'
        else:
            doc = Document(file)
            text = '\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
            record.update(native_text_chars=len(text), question_header_occurrences=len(re.findall(r'Câu\s*\d+', text, re.I)), embedded_images=len(doc.inline_shapes), strategy='native_docx_text_and_embedded_images')
        record['review_status'] = 'inventory_only_not_verified'
        documents.append(record)
    images = [str(p.relative_to(root)) for folder in ('DeThiPDF','MAT_MA') for p in (root/folder).rglob('*') if p.suffix.lower() in ('.png','.jpg','.jpeg')]
    result = {
        'scope': 'Existing bank: six source exams, five published. Other documents are inventory candidates, not silently merged or claimed verified.',
        'missing_original_directory': not (root/'DeThiOCR').is_dir(),
        'known_exams': [{**e, 'published': e['id'] in published, 'excluded_by_request': e['id']=='de-thi-k16-1'} for e in old],
        'documents': documents, 'standalone_scan_images': images,
        'counting_note': 'Question-header occurrences are NOT certified unique question counts; duplicated handouts/tables may repeat headers. Complete review required before import.',
    }
    target=WEB/'subjects/nmmmh/audit'; target.mkdir(parents=True,exist_ok=True)
    (target/'source_inventory.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'known_exams':len(old),'published_exams':len(published),'documents':len(documents),'scan_images':len(images),'missing_DeThiOCR':result['missing_original_directory']},ensure_ascii=False))


if __name__=='__main__': main()
